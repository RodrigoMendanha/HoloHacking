#!/usr/bin/env python3
"""
Versão editável (Word) do documento, montada a partir do HTML gerado por build.py.
Mesmo conteúdo; estilos do Word na identidade HoloHacking (Cormorant Garamond + Jost, verde e dourado).
Requer: pip install python-docx beautifulsoup4

  python3 scripts/doc-jornada/gerar_docx.py
"""
import pathlib, re
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RAIZ = pathlib.Path(__file__).resolve().parents[2]
DIR = RAIZ / 'docs' / 'jornada'
HTML = DIR / 'JORNADA-DA-NUTRICIONISTA-HOLOHACKING.html'
OUT = DIR / 'JORNADA-DA-NUTRICIONISTA-HOLOHACKING.docx'

VERDE = RGBColor(0x0B, 0x33, 0x25); VERDE_M = RGBColor(0x14, 0x50, 0x3B); DOURADO = RGBColor(0x8A, 0x6F, 0x32)
TINTA = RGBColor(0x1D, 0x3A, 0x2E); SUAVE = RGBColor(0x4F, 0x6A, 0x5D); CREME = RGBColor(0xF4, 0xEF, 0xE3)
FUNDO = {'pratica': 'FAF7EE', 'sistema': 'F1F5F1', 'decide': '0B3325', 'percebe': 'FBF6EA', 'atencao': 'FFFFFF',
         'resultado': 'EEF3EE', 'proximo': 'FFFFFF', 'futuro': 'FFF7E6', 'chave': 'FFFFFF', '': 'FAF7EE'}

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(20); sec.top_margin = Mm(20); sec.bottom_margin = Mm(20)

def fonte(estilo, nome, tam, cor=None, negrito=None):
    s = doc.styles[estilo]; s.font.name = nome; s.font.size = Pt(tam)
    s.element.rPr.rFonts.set(qn('w:eastAsia'), nome)
    if cor is not None: s.font.color.rgb = cor
    if negrito is not None: s.font.bold = negrito
fonte('Normal', 'Jost', 10, TINTA)
doc.styles['Normal'].paragraph_format.space_after = Pt(5); doc.styles['Normal'].paragraph_format.line_spacing = 1.25
fonte('Heading 1', 'Cormorant Garamond', 26, VERDE, True)
fonte('Heading 2', 'Cormorant Garamond', 17, VERDE, True)
fonte('Heading 3', 'Cormorant Garamond', 13.5, VERDE_M, True)
fonte('Title', 'Cormorant Garamond', 34, VERDE, True)
for h in ('Heading 1', 'Heading 2', 'Heading 3'):
    doc.styles[h].paragraph_format.space_before = Pt(12 if h != 'Heading 1' else 0)
    doc.styles[h].paragraph_format.keep_with_next = True

rodape = sec.footer.paragraphs[0]; rodape.text = 'HOLOHACKING  ·  Jornada da Nutricionista com a Paciente'
rodape.runs[0].font.size = Pt(7); rodape.runs[0].font.color.rgb = DOURADO

def sombrear(cel, hexa):
    tcPr = cel._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexa); tcPr.append(shd)

def borda_esquerda(cel, hexa='C9A35A', larg=18):
    tcPr = cel._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for lado in ('top', 'bottom', 'right'):
        e = OxmlElement('w:' + lado); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'E2DCCB'); b.append(e)
    e = OxmlElement('w:left'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(larg)); e.set(qn('w:color'), hexa); b.append(e)
    tcPr.append(b)

def runs(par, no, claro=False, base_bold=False, base_it=False):
    """escreve o texto inline (strong/em/br) de um nó num parágrafo"""
    for c in no.children:
        if isinstance(c, NavigableString):
            t = re.sub(r'\s+', ' ', str(c))
            if t.strip() or t == ' ':
                r = par.add_run(t); r.bold = base_bold or None; r.italic = base_it or None
                if claro: r.font.color.rgb = CREME
        elif isinstance(c, Tag):
            if c.name == 'br': par.add_run().add_break(); continue
            if c.name == 'img': continue
            runs(par, c, claro, base_bold or c.name in ('strong', 'b'), base_it or c.name in ('em', 'i'))

def paragrafo(destino, no, estilo=None, claro=False, tam=None, cor=None, centro=False, bold=False, it=False):
    p = destino.add_paragraph(style=estilo) if estilo else destino.add_paragraph()
    runs(p, no, claro, bold, it)
    for r in p.runs:
        if tam: r.font.size = Pt(tam)
        if cor is not None: r.font.color.rgb = cor
    if centro: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def lista(destino, no, claro=False):
    numerada = no.name == 'ol'
    for i, li in enumerate(no.find_all('li', recursive=False), 1):
        p = destino.add_paragraph(style='List Number' if numerada and destino is doc else 'List Bullet' if destino is doc else None)
        if destino is not doc: p.add_run(('%d. ' % i) if numerada else '• ')
        runs(p, li, claro)
        p.paragraph_format.space_after = Pt(2)

def imagem(destino, src, largura_mm):
    caminho = DIR / src
    if not caminho.exists(): return
    p = destino.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(caminho), width=Mm(largura_mm))

def tabela(no):
    linhas = no.find_all('tr')
    if not linhas: return
    ncol = max(len(l.find_all(['td', 'th'])) for l in linhas)
    t = doc.add_table(rows=0, cols=ncol); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, l in enumerate(linhas):
        cels = l.find_all(['td', 'th']); row = t.add_row().cells
        for j, c in enumerate(cels):
            p = row[j].paragraphs[0]; runs(p, c, claro=(c.name == 'th'))
            for r in p.runs: r.font.size = Pt(8.5)
            if c.name == 'th': sombrear(row[j], '0B3325'); [setattr(r, 'bold', True) for r in p.runs]
            elif k % 2 == 0: sombrear(row[j], 'FAF7EE')
    doc.add_paragraph()

def caixa(no, tipo, rotulo):
    t = doc.add_table(rows=1, cols=1); cel = t.cell(0, 0); sombrear(cel, FUNDO.get(tipo, 'FAF7EE')); borda_esquerda(cel)
    claro = tipo == 'decide'
    cel.paragraphs[0].text = ''
    if rotulo:
        r = cel.paragraphs[0].add_run(rotulo.upper()); r.bold = True; r.font.size = Pt(7.5); r.font.color.rgb = RGBColor(0xE2, 0xC9, 0x8A) if claro else DOURADO
    primeiro = not rotulo
    for c in no.children:
        if not isinstance(c, Tag) or 'box-rotulo' in (c.get('class') or []): continue
        alvo_p = None
        if c.name == 'p':
            alvo_p = cel.paragraphs[0] if primeiro else cel.add_paragraph(); runs(alvo_p, c, claro)
        elif c.name in ('ul', 'ol'):
            lista(cel, c, claro)
        elif c.name == 'div':
            alvo_p = cel.add_paragraph(); runs(alvo_p, c, claro)
        primeiro = False
        if tipo == 'chave' and alvo_p is not None:
            for r in alvo_p.runs: r.italic = True; r.font.size = Pt(13); r.font.name = 'Cormorant Garamond'; r.font.color.rgb = VERDE_M
    doc.add_paragraph()

def definicoes(titulo, subtitulo, dl, fundo='14503B'):
    t = doc.add_table(rows=1, cols=2); t.style = 'Table Grid'
    cab = t.cell(0, 0).merge(t.cell(0, 1)); sombrear(cab, fundo)
    p = cab.paragraphs[0]; r = p.add_run(titulo); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = CREME; r.font.name = 'Cormorant Garamond'
    if subtitulo:
        r = p.add_run('   ' + subtitulo.upper()); r.font.size = Pt(7.5); r.font.color.rgb = RGBColor(0xE2, 0xC9, 0x8A)
    for dt in dl.find_all('dt'):
        dd = dt.find_next_sibling('dd'); row = t.add_row().cells
        p = row[0].paragraphs[0]; r = p.add_run(dt.get_text(' ', strip=True).upper()); r.bold = True; r.font.size = Pt(7.5); r.font.color.rgb = DOURADO
        p = row[1].paragraphs[0]; runs(p, dd)
        for r in p.runs: r.font.size = Pt(9)
    for row in t.rows: row.cells[0].width = Mm(40); row.cells[1].width = Mm(130)
    doc.add_paragraph()

def quebra(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

sopa = BeautifulSoup(HTML.read_text(encoding='utf-8'), 'html.parser')
primeiro_h1 = True
for no in sopa.body.children:
    if not isinstance(no, Tag): continue
    cls = no.get('class') or []
    if 'capa' in cls:
        doc.add_paragraph(); imagem(doc, 'imagens/logo-pinha.png', 45)
        p = doc.add_paragraph(style='Title'); p.add_run(no.find('h1').get_text(strip=True)); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for s in no.select('.sub, .sub2, .rodape-capa'):
            q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = q.add_run(s.get_text(' ', strip=True))
            r.italic = 'sub' in s.get('class'); r.font.color.rgb = DOURADO if 'sub' in s.get('class') else SUAVE
            r.font.size = Pt(14 if s.get('class') == ['sub'] else 9)
        quebra(); continue
    if no.name == 'h1':
        if not primeiro_h1: quebra()
        primeiro_h1 = False
        doc.add_paragraph(no.get_text(strip=True), style='Heading 1'); continue
    if no.name in ('h2', 'h3'):
        doc.add_paragraph(no.get_text(strip=True), style='Heading 2' if no.name == 'h2' else 'Heading 3'); continue
    if no.name == 'p':
        if 'abre' in cls: paragrafo(doc, no, tam=13, cor=SUAVE, it=True)
        elif 'eyebrow' in cls: paragrafo(doc, no, tam=8, cor=DOURADO, bold=True)
        else: paragrafo(doc, no)
        continue
    if no.name in ('ul', 'ol'): lista(doc, no); continue
    if no.name == 'table': tabela(no); continue
    if no.name == 'ol' and 'sumario' in cls: lista(doc, no); continue
    if no.name == 'figure':
        img = no.find('img'); largura = 120 if 'estreita' in (img.get('class') or []) else 150 if 'media' in (img.get('class') or []) else 170
        if 'estreita' in (img.get('class') or []): largura = 75
        imagem(doc, img['src'], largura)
        cap = no.find('figcaption')
        if cap: paragrafo(doc, cap, tam=8, cor=SUAVE)
        continue
    if no.name == 'div':
        if 'box' in cls:
            tipo = next((c for c in cls if c in FUNDO and c), '')
            rot = no.find('p', class_='box-rotulo'); caixa(no, tipo, rot.get_text(strip=True) if rot else ''); continue
        if 'frase' in cls:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            small = no.find('small'); txt = no.find('p') or no
            runs(p, txt)
            for r in p.runs: r.font.name = 'Cormorant Garamond'; r.font.size = Pt(15); r.font.color.rgb = VERDE
            if small:
                q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = q.add_run(small.get_text(strip=True).upper()); r.font.size = Pt(7.5); r.font.color.rgb = DOURADO
            continue
        if any(c in cls for c in ('duas', 'tres')):
            for filho in no.find_all(['div', 'figure'], recursive=False):
                fc = filho.get('class') or []
                if filho.name == 'figure':
                    img = filho.find('img'); imagem(doc, img['src'], 85); cap = filho.find('figcaption'); cap and paragrafo(doc, cap, tam=8, cor=SUAVE)
                elif 'box' in fc:
                    tipo = next((c for c in fc if c in FUNDO and c), ''); rot = filho.find('p', class_='box-rotulo'); caixa(filho, tipo, rot.get_text(strip=True) if rot else '')
                else:
                    for x in filho.find_all(['p', 'ul', 'ol', 'figure'], recursive=False):
                        if x.name == 'p': paragrafo(doc, x)
                        elif x.name == 'figure':
                            img = x.find('img'); imagem(doc, img['src'], 80); cap = x.find('figcaption'); cap and paragrafo(doc, cap, tam=8, cor=SUAVE)
                        else: lista(doc, x)
                    for b in filho.find_all('div', class_='box', recursive=False):
                        rot = b.find('p', class_='box-rotulo'); caixa(b, next((c for c in b['class'] if c in FUNDO and c), ''), rot.get_text(strip=True) if rot else '')
            continue
        if any(c in cls for c in ('fluxo', 'linha-tempo', 'principios', 'papeis', 'cadeia')):
            itens = no.find_all(['div', 'span'], recursive=False)
            if 'cadeia' in cls:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run('  →  '.join(s.get_text(strip=True) for s in no.find_all('span'))); r.bold = True; r.font.color.rgb = VERDE_M; r.font.size = Pt(9)
                continue
            for k, it in enumerate(itens, 1):
                p = doc.add_paragraph(style='List Bullet')
                b = it.find('b')
                if 'linha-tempo' in cls: p.add_run('Atendimento %s — ' % it.get('data-n')).bold = True
                if b: p.add_run(b.get_text(strip=True) + ': ').bold = True
                resto = it.get_text(' ', strip=True); resto = resto[len(b.get_text(strip=True)):].strip() if b else resto
                p.add_run(resto)
            continue
        if 'ferr' in cls:
            cab = no.find('div', class_='ferr-cab')
            fundo = '14503B' if 'corpo' in cls else '0B3325' if 'mente' in cls else '06231A'
            definicoes(cab.find('b').get_text(strip=True), cab.find('span').get_text(strip=True), no.find('dl'), fundo); continue
        if 'encontro' in cls:
            cab = no.find('div', class_='encontro-cab')
            definicoes('Atendimento ' + cab.find(class_='n').get_text(strip=True) + ' — ' + cab.find('b').get_text(strip=True), cab.find('span', class_=None).get_text(strip=True) if cab.find('span', class_=None) else '', no.find('dl'), '0B3325'); continue
        if 'fim' in cls: imagem(doc, 'imagens/logo-pinha.png', 25); continue
        if 'quebra' in cls: quebra(); continue
        # outros blocos: texto simples
        for x in no.find_all(['p', 'ul', 'ol', 'table'], recursive=False):
            if x.name == 'p': paragrafo(doc, x)
            elif x.name == 'table': tabela(x)
            else: lista(doc, x)

doc.core_properties.title = 'HoloHacking — Jornada da Nutricionista com a Paciente'
doc.core_properties.subject = 'Primeiro modelo · outubro de 2026'
doc.save(OUT)
print('docx', OUT.relative_to(RAIZ), OUT.stat().st_size // 1024, 'KB')
